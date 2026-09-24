import io
import json
import zipfile

import pytest
from PIL import Image

from netra_ocr.document import AssetStore, Document, DocumentValidationError
from netra_ocr.exporters import export


def make_doc():
    assets = AssetStore()
    assets.put("page-0", Image.new("RGB", (600, 800), "white"))
    assets.put("page-1", Image.new("RGB", (600, 800), "white"))
    assets.put("p0-img3", Image.new("RGB", (120, 80), "red"))
    data = {
        "pages": [{"width": 600, "height": 800, "image": "page-0"},
                  {"width": 600, "height": 800, "image": "page-1"}],
        "blocks": [
            {"id": "h", "type": "page_header", "text": "ក្រសួង", "page": 0},
            {"id": "a", "type": "heading", "level": 1, "text": "ព្រះរាជាណាចក្រកម្ពុជា\nជាតិ សាសនា", "align": "center", "page": 0},
            {"id": "b", "type": "paragraph", "text": "Body | text", "page": 0, "bbox": [0, 0, 600, 40]},
            {"id": "c", "type": "figure", "image": "p0-img3", "page": 0, "bbox": [0, 100, 120, 180]},
            {"id": "d", "type": "caption", "text": "Figure 1", "page": 0},
            {"id": "e", "type": "table", "rows": [["ឈ្មោះ", "អាយុ"], ["Dara", "31"]], "page": 1},
            {"id": "f", "type": "page_footer", "text": "p. 2", "page": 1},
        ],
    }
    return Document.from_dict(data, assets)


def test_validation_rejects_bad_input():
    doc = make_doc()
    base = doc.to_dict()
    for mutate in (
        lambda d: d["blocks"].append({"id": "a", "type": "paragraph", "text": "dup id", "page": 0}),
        lambda d: d["blocks"].append({"id": "z", "type": "script", "text": "x", "page": 0}),
        lambda d: d["blocks"].append({"id": "z", "type": "figure", "image": "../../etc/passwd", "page": 0}),
        lambda d: d["blocks"].append({"id": "z", "type": "paragraph", "text": "x", "page": 9}),
        lambda d: d["blocks"].append({"id": "z", "type": "heading", "level": 7, "text": "x", "page": 0}),
        lambda d: d["blocks"].append({"id": "z", "type": "table", "rows": [], "page": 0}),
    ):
        bad = json.loads(json.dumps(base))
        mutate(bad)
        with pytest.raises(DocumentValidationError):
            Document.from_dict(bad, doc.assets)


def test_save_load_roundtrip(tmp_path):
    doc = make_doc()
    doc.save(str(tmp_path))
    again = Document.load(str(tmp_path))
    assert again.to_dict() == doc.to_dict()
    assert again.assets.get("p0-img3").size == (120, 80)


def test_all_formats_export():
    doc = make_doc()
    for fmt in ("docx", "html", "md", "json", "txt"):
        data, name, mime = export(doc, fmt)
        assert data and name.startswith("document")
    html = export(doc, "html")[0].decode()
    assert "<h1 style=\"text-align:center\">" in html and "<th>ឈ្មោះ</th>" in html
    assert "Body | text" in html and "data:image/png;base64," in html
    md_zip = zipfile.ZipFile(io.BytesIO(export(doc, "md")[0]))
    md = md_zip.read("document.md").decode()
    assert "# ព្រះរាជាណាចក្រកម្ពុជា — ជាតិ សាសនា" in md
    assert "| ឈ្មោះ | អាយុ |" in md and "---" in md and "images/p0-img3.png" in md_zip.namelist()
    assert json.loads(export(doc, "json")[0])["blocks"][1]["level"] == 1
    assert "Dara\t31" in export(doc, "txt")[0].decode()


def test_docx_is_structured():
    from docx import Document as Word
    from docx.oxml.ns import qn
    word = Word(io.BytesIO(export(make_doc(), "docx")[0]))
    styles = [(p.style.name, p.text) for p in word.paragraphs if p.text]
    assert ("Heading 1", "ព្រះរាជាណាចក្រកម្ពុជា\nជាតិ សាសនា") in styles
    assert ("Caption", "Figure 1") in styles
    assert len(word.tables) == 1
    t = word.tables[0]
    assert [[c.text for c in r.cells] for r in t.rows] == [["ឈ្មោះ", "អាយុ"], ["Dara", "31"]]
    assert len(word.sections) == 2                            # one section (page) per source page
    assert word.sections[0].header.paragraphs[0].text == "ក្រសួង"
    assert word.sections[1].footer.paragraphs[0].text == "p. 2"
    rfonts = word.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:cs")) == "Kantumruy Pro"          # complex-script (Khmer) font
    assert len(word.inline_shapes) == 1
