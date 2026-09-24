"""Structured Word export: real headings, paragraphs, tables, figures and captions.

Each source page becomes its own Word section (new page), so a page's
detected header/footer text lands in that section's real header/footer.
Khmer is a complex script in OOXML: fonts, sizes and bold have to be set on
the complex-script (``cs``) attributes too, not just ascii/hAnsi, or Word
renders Khmer runs in its fallback font.
"""

import io

from ..document import Document

FONT = "Kantumruy Pro"
INK = "191014"
MAROON = "4f2724"
MUTED = "6b6266"
BORDER = "e7e1e0"
WASH = "f7f5f5"        # brand --wash (maroon at 4.5%) flattened onto white


def _import_docx():
    try:
        import docx  # noqa: F401
    except ImportError as e:
        raise ImportError(
            "Writing .docx output requires the 'docx' extra. Install it with: "
            "pip install \"netra-ocr[docx]\""
        ) from e


def _set_style_font(style, size_pt: float, bold: bool = False, color: str = INK, italic: bool = False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    font = style.font
    font.name = FONT
    font.size = Pt(size_pt)
    font.bold = bold
    font.italic = italic
    font.color.rgb = RGBColor.from_string(color.upper())
    rpr = style.element.get_or_add_rPr()

    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(theme_attr)) is not None:
            del rfonts.attrib[qn(theme_attr)]

    def _ensure(tag, **attrs):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
        return el

    _ensure("w:szCs", **{"w:val": str(int(size_pt * 2))})
    if bold:
        _ensure("w:bCs")
    _ensure("w:lang", **{"w:val": "en-US", "w:bidi": "km-KH"})


def _configure_styles(doc):
    from docx.shared import Pt

    styles = doc.styles
    _set_style_font(styles["Normal"], 11)
    styles["Normal"].paragraph_format.line_spacing = 1.3
    styles["Normal"].paragraph_format.space_after = Pt(6)
    _set_style_font(styles["Heading 1"], 18, bold=True)
    _set_style_font(styles["Heading 2"], 14, bold=True, color=MAROON)
    _set_style_font(styles["Heading 3"], 12, bold=True, color=MAROON)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        pf = styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(12), Pt(6)
    _set_style_font(styles["Caption"], 9.5, italic=True, color=MUTED)
    _set_style_font(styles["Header"], 9, color=MUTED)
    _set_style_font(styles["Footer"], 9, color=MUTED)


def _align(value: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    return {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(value, WD_ALIGN_PARAGRAPH.LEFT)


def _write_text(paragraph, text: str, bold: bool = False):
    """Add ``text`` to a paragraph, turning "\n" into line breaks."""
    for i, part in enumerate(text.split("\n")):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        if bold:
            from docx.oxml import OxmlElement
            run.bold = True
            run._r.get_or_add_rPr().append(OxmlElement("w:bCs"))


def _shade(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _table_borders(table, color: str = BORDER):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def _picture(paragraph, doc: Document, asset_id: str, bbox, page_width_px: int, max_width_in: float):
    """Insert an asset, scaled to the share of the page width it had in the scan."""
    from docx.shared import Inches
    png = doc.assets.png_bytes(asset_id)
    if png is None:
        return
    frac = 1.0
    if bbox and page_width_px:
        frac = max(0.08, min(1.0, (bbox[2] - bbox[0]) / page_width_px))
    paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(max_width_in * frac))


def to_docx(doc: Document) -> bytes:
    _import_docx()
    from docx import Document as WordDocument
    from docx.enum.section import WD_SECTION
    from docx.shared import Cm

    word = WordDocument()
    _configure_styles(word)
    usable_width_in = (21.0 - 2 * 2.0) / 2.54   # A4, 2 cm margins

    def setup_section(section):
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.0))

    page_widths = {p["index"]: p.get("width") for p in doc.pages}
    by_page: dict[int, list] = {}
    for b in doc.blocks:
        by_page.setdefault(b["page"], []).append(b)
    pages = sorted(by_page) or [0]

    first = True
    for page_index in pages:
        section = word.sections[0] if first else word.add_section(WD_SECTION.NEW_PAGE)
        setup_section(section)
        if not first:
            section.header.is_linked_to_previous = False
            section.footer.is_linked_to_previous = False
        header_used = footer_used = False
        first = False
        page_w = page_widths.get(page_index) or 0

        for b in by_page.get(page_index, []):
            t = b["type"]
            if t in ("page_header", "page_footer"):
                part = section.header if t == "page_header" else section.footer
                used = header_used if t == "page_header" else footer_used
                para = part.paragraphs[0] if not used and part.paragraphs else part.add_paragraph()
                para.style = word.styles["Header" if t == "page_header" else "Footer"]
                para.alignment = _align(b.get("align", "left"))
                _write_text(para, b["text"])
                if t == "page_header":
                    header_used = True
                else:
                    footer_used = True
            elif t == "heading":
                para = word.add_paragraph(style=f"Heading {b.get('level', 1)}")
                para.alignment = _align(b.get("align", "left"))
                _write_text(para, b["text"])
            elif t == "paragraph":
                para = word.add_paragraph(style="Normal")
                para.alignment = _align(b.get("align", "left"))
                _write_text(para, b["text"])
            elif t == "caption":
                para = word.add_paragraph(style="Caption")
                para.alignment = _align(b.get("align", "left"))
                _write_text(para, b["text"])
            elif t in ("figure", "formula"):
                para = word.add_paragraph()
                para.alignment = _align("center")
                _picture(para, doc, b["image"], b.get("bbox"), page_w, usable_width_in)
            elif t == "table":
                rows = b["rows"]
                if len(rows) == 1 and len(rows[0]) == 1 and b.get("image"):
                    # No grid could be recovered: keep the original table
                    # image so no structure is lost, with the text below it.
                    para = word.add_paragraph()
                    _picture(para, doc, b["image"], b.get("bbox"), page_w, usable_width_in)
                ncols = max(len(r) for r in rows)
                table = word.add_table(rows=len(rows), cols=ncols)
                _table_borders(table)
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        cell = table.cell(ri, ci)
                        cell.paragraphs[0].style = word.styles["Normal"]
                        header = ri == 0 and len(rows) > 1
                        _write_text(cell.paragraphs[0], row[ci] if ci < len(row) else "", bold=header)
                        if header:
                            _shade(cell, WASH)
                if b.get("caption"):
                    word.add_paragraph(b["caption"], style="Caption")
                word.add_paragraph()

    buf = io.BytesIO()
    word.save(buf)
    return buf.getvalue()
